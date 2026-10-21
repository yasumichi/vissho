"""Tests for Excel exporter."""

from pathlib import Path
from openpyxl import load_workbook
import pytest

from vissho.config import ColumnDefinition, Prerequisites, VisshoConfig
from vissho.core.schema import create_curriculum_plan_model
from vissho.core.validator import validate_curriculum_plan
from vissho.exporter.excel import export_to_excel


@pytest.fixture
def sample_config() -> VisshoConfig:
    return VisshoConfig(
        prerequisites=Prerequisites(
            purpose="テスト講座（Python基礎）",
            target_audience="専門学校1年生",
            prerequisite_knowledge="PC基本操作",
            total_sessions=2,
            session_minutes=90,
        ),
        columns=[
            ColumnDefinition(name="session_number", description="コマ番号", width=12),
            ColumnDefinition(name="topic", description="単元・テーマ", width=25),
            ColumnDefinition(name="allocated_minutes", description="配当時間", width=14),
            ColumnDefinition(name="activities", description="学習活動", width=40),
            ColumnDefinition(name="source_slides", description="出典スライド", width=16),
        ],
    )


@pytest.fixture
def sample_plan_data() -> dict:
    return {
        "sessions": [
            {
                "session_number": 1,
                "session_title": "Python導入",
                "rows": [
                    {
                        "session_number": "第1コマ",
                        "topic": "ガイダンスと環境設定",
                        "allocated_minutes": 45,
                        "activities": "ツールのインストール\n動作確認",
                        "source_slides": [1, 2],
                    },
                    {
                        "session_number": "第1コマ",
                        "topic": "変数と型",
                        "allocated_minutes": 45,
                        "activities": "数値と文字列の演算",
                        "source_slides": [3],
                    },
                ],
                "total_minutes": 90,
            },
            {
                "session_number": 2,
                "session_title": "条件分岐とループ",
                "rows": [
                    {
                        "session_number": "第2コマ",
                        "topic": "if文による分岐",
                        "allocated_minutes": 90,
                        "activities": "演習課題の作成",
                        "source_slides": [4, 5],
                    }
                ],
                "total_minutes": 90,
            },
        ],
        "unassigned_slides": [6],
    }


def test_export_to_excel_structure(sample_config: VisshoConfig, sample_plan_data: dict, tmp_path: Path) -> None:
    """Test Excel workbook creation, structure, sheet name, and column widths."""
    target_file = tmp_path / "plan.xlsx"
    wb = export_to_excel(sample_plan_data, sample_config, output_path=target_file)

    assert target_file.exists()
    assert "学習指導案" in wb.sheetnames
    ws = wb["学習指導案"]

    # Verify column widths match config metadata
    for col_idx, col in enumerate(sample_config.columns, start=1):
        from openpyxl.utils import get_column_letter

        col_letter = get_column_letter(col_idx)
        assert ws.column_dimensions[col_letter].width == float(col.width)


def test_export_to_excel_no_bold_no_fill_on_headers(sample_config: VisshoConfig, sample_plan_data: dict) -> None:
    """Strictly verify specification: NO background fill and NO bold font on table headers."""
    wb = export_to_excel(sample_plan_data, sample_config)
    ws = wb["学習指導案"]

    # Search for header row containing column names
    header_found = False
    for row in ws.iter_rows():
        values = [c.value for c in row if c.value is not None]
        if "topic" in values and "allocated_minutes" in values:
            header_found = True
            for c in row:
                if c.value:
                    # Must NOT have bold font
                    assert c.font is not None
                    assert c.font.bold is False, f"Header cell {c.coordinate} has bold=True"
                    # Must NOT have background fill
                    assert c.fill.fill_type is None, f"Header cell {c.coordinate} has fill_type={c.fill.fill_type}"

    assert header_found, "Table header row was not found in worksheet"


def test_export_to_excel_wrap_text_and_formulas(sample_config: VisshoConfig, sample_plan_data: dict, tmp_path: Path) -> None:
    """Test text wrapping on cells and SUM formula generation for session totals."""
    target_file = tmp_path / "plan_formula.xlsx"
    export_to_excel(sample_plan_data, sample_config, output_path=target_file)

    # Re-open workbook to inspect stored formula and cell attributes
    loaded_wb = load_workbook(target_file, data_only=False)
    ws = loaded_wb["学習指導案"]

    # Find SUM formulas in allocated_minutes column
    found_session_sums = []
    grand_sum_found = False

    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and isinstance(cell.value, str):
                if cell.value.startswith("=SUM("):
                    found_session_sums.append(cell.value)
                elif cell.value.startswith("=") and "+" in cell.value:
                    grand_sum_found = True

    # 2 sessions should each have a SUM formula
    assert len(found_session_sums) == 2
    assert grand_sum_found is True

    # Verify wrap text is enabled on data cells
    multiline_cell_found = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and "ツールのインストール" in str(cell.value):
                multiline_cell_found = True
                assert cell.alignment.wrap_text is True

    assert multiline_cell_found is True


def test_export_to_excel_with_validation_report(sample_config: VisshoConfig, sample_plan_data: dict, tmp_path: Path) -> None:
    """Test Excel export includes validation sheet when validation report is provided."""
    plan_cls = create_curriculum_plan_model(sample_config.columns)
    model_plan = plan_cls.model_validate(sample_plan_data)

    report = validate_curriculum_plan(model_plan, sample_config, all_slide_numbers=[1, 2, 3, 4, 5, 6])

    target_file = tmp_path / "plan_validated.xlsx"
    wb = export_to_excel(model_plan, sample_config, output_path=target_file, validation_report=report)

    assert "検証結果" in wb.sheetnames
    ws_val = wb["検証結果"]

    # Check validation report content
    cell_values = [str(c.value) for row in ws_val.iter_rows() for c in row if c.value is not None]
    joined_text = " ".join(cell_values)

    assert "PASSED" in joined_text or "合格" in joined_text
    assert "第1コマ" in joined_text
    assert "第2コマ" in joined_text


def test_export_to_excel_without_explicit_allocated_minutes_col(tmp_path: Path) -> None:
    """Test Excel export adds allocated_minutes column automatically if omitted in config."""
    config = VisshoConfig(
        prerequisites=Prerequisites(
            purpose="テスト講座",
            target_audience="専門生",
            total_sessions=1,
            session_minutes=90,
        ),
        columns=[
            ColumnDefinition(name="topic", description="単元", width=25),
        ],
    )
    plan_data = {
        "sessions": [
            {
                "session_number": 1,
                "session_title": "単元1",
                "rows": [
                    {
                        "topic": "導入",
                        "allocated_minutes": 90,
                    }
                ],
                "total_minutes": 90,
            }
        ],
        "unassigned_slides": [],
    }
    target_file = tmp_path / "auto_col.xlsx"
    wb = export_to_excel(plan_data, config, output_path=target_file)
    assert target_file.exists()
    ws = wb["学習指導案"]

    header_values = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert "allocated_minutes" in header_values


def test_export_to_excel_with_validation_errors(sample_config: VisshoConfig, tmp_path: Path) -> None:
    """Test Excel validation sheet when validation report has errors and warnings."""
    from vissho.core.validator import SessionValidation, ValidationReport

    report = ValidationReport(
        is_valid=False,
        total_sessions_valid=False,
        total_sessions_actual=1,
        total_sessions_target=2,
        total_minutes_actual=45,
        total_minutes_target=180,
        total_minutes_difference=-135,
        session_validations=[
            SessionValidation(
                session_number=1,
                session_title="不完全コマ",
                allocated_minutes=45,
                target_minutes=90,
                difference=-45,
                is_valid=False,
            )
        ],
        referenced_slides=[1],
        unassigned_slides=[2],
        missing_slides=[3],
        invalid_slide_references=[99],
        errors=["コマ数不一致: 目標2コマに対して1コマです。"],
        warnings=["スライド99が存在しません。"],
    )

    plan_data = {
        "sessions": [
            {
                "session_number": 1,
                "session_title": "不完全コマ",
                "rows": [
                    {
                        "session_number": "第1コマ",
                        "topic": "導入",
                        "allocated_minutes": 45,
                        "source_slides": [1],
                    }
                ],
                "total_minutes": 45,
            }
        ],
        "unassigned_slides": [2],
    }

    target_file = tmp_path / "errors.xlsx"
    wb = export_to_excel(plan_data, sample_config, output_path=target_file, validation_report=report)
    assert "検証結果" in wb.sheetnames
    ws_val = wb["検証結果"]
    text = " ".join(str(c.value) for row in ws_val.iter_rows() for c in row if c.value)
    assert "ISSUES DETECTED" in text or "指摘事項あり" in text
    assert "エラー一覧:" in text
    assert "警告一覧:" in text
    assert "未言及（欠落）スライド:" in text
    assert "不正なスライド参照:" in text
